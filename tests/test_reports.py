from openpyxl import load_workbook

from leadcore import pipeline
from leadcore.report import capture_sheet, tracker
from leadcore.report.capture_sheet import excel_template

from .conftest import ListLoader, rec


def test_excel_template_uses_brace_tokens():
    assert excel_template("first.last") == "{F}.{L}"
    assert excel_template("flast") == "{f}{L}"
    assert excel_template("first") == "{F}"
    assert excel_template(None) == ""


def _run(session):
    pipeline.ingest(session, loaders=[ListLoader([
        rec("jane.doe@a.com", "Jane", row=2), rec("john.smith@a.com", "John", row=3),
        rec("mary@b.com", "Mary", row=4), rec("pjones@b.com", "Paul", row=5),
        rec("alex.lee@c.com", "there", row=6),
    ])])
    return pipeline.run_inference(session, use_nicknames=True)[0]


def test_tracker_lists_every_firm_problems_first(session, tmp_path):
    path = tracker.build(session, _run(session), tmp_path / "t.xlsx")
    rows = list(load_workbook(path)["Firms"].iter_rows(min_row=2, values_only=True))
    assert [(r[0], r[1]) for r in rows] == [("b.com", "D"), ("c.com", "E"), ("a.com", "B")]
    assert "mary@b.com" in rows[0][9] and "pjones@b.com" in rows[0][9]
    assert rows[1][11] == "alex.lee@c.com"


def test_capture_sheet_has_prediction_formula(session, tmp_path):
    path = capture_sheet.build(session, _run(session), per_tier=5, seed=1, path=tmp_path / "v.xlsx")
    ws = load_workbook(path)["Validation"]
    assert str(ws["I2"].value).startswith("=IF(") and ws["N2"].value == "{F}.{L}"
