"""Pattern workbook: data/out/patterns.xlsx (Summary, Domain patterns, Tier D conflicts, Excluded)."""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from leadcore import config
from leadcore.models import Contact, DomainPattern, EmailObservation, InferenceRun, Origin
from leadcore.patterns import library

HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TIER_FILLS = {t: PatternFill("solid", start_color=c) for t, c in
              {"A": "C6EFCE", "B": "DDEBF7", "C": "FFF2CC", "D": "F8CBAD", "E": "EDEDED"}.items()}


def _header(ws, cols: list[str], widths: list[int]) -> None:
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def worked_example(pattern: str | None, domain: str) -> str:
    return library.get(pattern).example(domain) if pattern else ""


def _seen(d: dict) -> str:
    return ", ".join(f"{p} ({n})" for p, n in d.items())


def overlap_stats(session: Session) -> dict[str, int]:
    """Business-contact overlap between the two source files (post-exclusion)."""
    rows = session.execute(
        select(EmailObservation.email, EmailObservation.source)
        .where(EmailObservation.email.is_not(None),
               EmailObservation.excluded_reason.is_(None) | (EmailObservation.excluded_reason == "placeholder_name"))
        .distinct()
    ).all()
    by: dict[str, set] = {}
    for e, s in rows:
        by.setdefault(e, set()).add(s)
    raw = session.execute(
        select(EmailObservation.email, EmailObservation.source)
        .where(EmailObservation.email.is_not(None)).distinct()
    ).all()
    by_raw: dict[str, set] = {}
    for e, s in raw:
        by_raw.setdefault(e, set()).add(s)
    return {
        "business": sum(1 for s in by.values() if len(s) > 1),
        "all_valid": sum(1 for s in by_raw.values() if len(s) > 1),
    }


def summary_rows(session: Session, run: InferenceRun) -> list[tuple]:
    obs_total = session.scalar(select(func.count()).select_from(EmailObservation))
    by_source = dict(session.execute(select(EmailObservation.source, func.count())
                                     .group_by(EmailObservation.source)).all())
    contacts = session.scalar(select(func.count()).select_from(Contact).where(Contact.origin == Origin.confirmed))
    generated = session.scalar(select(func.count()).select_from(Contact).where(Contact.origin == Origin.generated))
    sendable = session.scalar(select(func.count()).select_from(Contact.sendable().subquery()))
    tiers = Counter(dict(session.execute(select(DomainPattern.tier, func.count())
                                         .where(DomainPattern.run_id == run.id)
                                         .group_by(DomainPattern.tier)).all()))
    dominant = Counter(dict(session.execute(
        select(DomainPattern.pattern, func.count())
        .where(DomainPattern.run_id == run.id, DomainPattern.tier.in_(config.GENERATABLE_TIERS))
        .group_by(DomainPattern.pattern)).all()))
    ov = overlap_stats(session)
    rows: list[tuple] = [
        ("Corpus", None),
        ("Observations stored", obs_total),
        *[(f"  from {k}", v) for k, v in sorted(by_source.items())],
        ("Unique business contacts (confirmed, post-exclusion)", contacts),
        ("Unique company domains", sum(tiers.values())),
        ("Email overlap between files (business contacts)", ov["business"]),
        ("Email overlap between files (all valid addresses)", ov["all_valid"]),
        ("Generated contacts held", generated),
        ("Sendable via Contact.sendable()", sendable),
        (None, None),
        ("Inference run", f"#{run.id}  nicknames={'on' if run.use_nicknames else 'off'}  as_of={run.config.get('as_of')}"),
        *[(f"Tier {t}", tiers.get(t, 0)) for t in "ABCDE"],
        ("Generatable (A/B/C)", sum(tiers.get(t, 0) for t in config.GENERATABLE_TIERS)),
        (None, None),
        ("Dominant conventions (Tier A-C domains)", None),
        *[(f"  {p}  e.g. {library.get(p).example()}", n) for p, n in dominant.most_common(12) if p],
    ]
    return rows


def build(session: Session, run: InferenceRun, path: Path | None = None) -> Path:
    path = Path(path or config.DATA_OUT / "patterns.xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    _header(ws, ["Metric", "Value"], [58, 40])
    for k, v in summary_rows(session, run):
        ws.append([k, v])
        if k and v is None:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    dps = session.scalars(select(DomainPattern).where(DomainPattern.run_id == run.id)
                          .order_by(DomainPattern.tier, DomainPattern.sample_count.desc(), DomainPattern.domain)).all()

    ws = wb.create_sheet("Domain patterns")
    cols = ["Domain", "Tier", "Inferred pattern", "Worked example", "Samples", "Unexplained",
            "Dominance", "All patterns seen (count)", "Contacts held", "Evidence address", "Note"]
    _header(ws, cols, [30, 6, 16, 38, 9, 11, 10, 46, 9, 40, 26])
    for dp in dps:
        note = dp.note or ""
        if dp.migrated_from:
            note = f"migrated from {dp.migrated_from}"
        ws.append([dp.domain, dp.tier, dp.pattern if dp.tier in "ABC" else (dp.pattern and f"({dp.pattern})"),
                   worked_example(dp.pattern, dp.domain) if dp.tier in "ABC" else "",
                   dp.sample_count, dp.unexplained_count, dp.dominance, _seen(dp.patterns_seen),
                   dp.contacts_held, dp.evidence_email, note])
        ws.cell(row=ws.max_row, column=2).fill = TIER_FILLS[dp.tier]
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    ws = wb.create_sheet("Tier D conflicts")
    cols = ["Domain", "Samples", "Unexplained", "Leading pattern", "Leading share", "Competing patterns (count)",
            "Contacts held"]
    _header(ws, cols, [30, 9, 11, 16, 12, 60, 9])
    for dp in dps:
        if dp.tier == "D":
            ws.append([dp.domain, dp.sample_count, dp.unexplained_count, dp.pattern, dp.dominance,
                       _seen(dp.patterns_seen), dp.contacts_held])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    ws = wb.create_sheet("Excluded")
    _header(ws, ["Reason", "Source file", "Observations", "Distinct addresses"], [22, 28, 14, 18])
    rows = session.execute(
        select(EmailObservation.excluded_reason, EmailObservation.source_file, func.count(),
               func.count(func.distinct(EmailObservation.email)))
        .where(EmailObservation.excluded_reason.is_not(None))
        .group_by(EmailObservation.excluded_reason, EmailObservation.source_file)
        .order_by(EmailObservation.excluded_reason, EmailObservation.source_file)).all()
    for r in rows:
        ws.append(list(r))
    ws.append([])
    ws.append(["Note: placeholder_name rows remain business contacts (their address is real); "
               "they are excluded only from pattern inference."])
    ws.append([])
    ws.append(["Examples"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for reason in ("invalid_email", "consumer_domain", "role_address", "placeholder_name"):
        ex = session.execute(
            select(EmailObservation.email_raw, EmailObservation.first_name_raw, EmailObservation.source_file)
            .where(EmailObservation.excluded_reason == reason).order_by(EmailObservation.id).limit(5)).all()
        for e in ex:
            ws.append([reason, e[2], e[0], e[1]])

    wb.save(path)
    return path
