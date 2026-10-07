"""Firm tracker: data/out/firm_tracker.xlsx.

One sheet listing every firm, worst problems first, with the evidence a person
needs to resolve it and yellow columns for the answer. Built for handing to
the sales team: conflicting firms (Tier D) show which named people use which
format, so someone who knows the firm can say which one is right.
"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from leadcore import config, pipeline
from leadcore.ingest.exclusions import is_placeholder_name
from leadcore.models import DomainPattern, EmailObservation, InferenceRun
from leadcore.patterns import inference, library
from leadcore.report.capture_sheet import INPUT_FILL
from leadcore.report.workbook import HEADER_FILL, HEADER_FONT, TIER_FILLS, worked_example

# (status, action needed) per tier, in priority order.
STATUS = {
    "D": ("Conflicting formats", "Which format does this firm use for most staff? Pick it in 'Correct format'."),
    "E": ("No usable names", "Give one known person's full name and work address at this firm."),
    "C": ("One address only", "Confirm with one more known address, or confirm the format."),
    "B": ("Two agreeing addresses", "Optional spot-check."),
    "A": ("Confirmed", "None: ready to generate."),
}
PRIORITY = {t: i for i, t in enumerate(STATUS)}
EXTRA_CHOICES = ("mixed / depends on person", "unknown")
EVIDENCE_PER_FORMAT = 3


def _evidence(samples: list[inference.Sample]) -> tuple[str, str]:
    """('format: Name <addr>; ...' grouped by format, 'Name <addr>; ...' for unexplained)."""
    by_format: dict[str, list[str]] = defaultdict(list)
    unexplained = []
    for s in sorted(samples, key=lambda x: (-x.weight, x.email)):
        item = f"{s.first_name} <{s.email}>"
        if s.patterns:
            by_format[s.patterns[0]].append(item)
        else:
            unexplained.append(item)
    parts = [f"{p}: " + "; ".join(v[:EVIDENCE_PER_FORMAT]) + (f" (+{len(v) - EVIDENCE_PER_FORMAT})"
             if len(v) > EVIDENCE_PER_FORMAT else "")
             for p, v in sorted(by_format.items(), key=lambda kv: -len(kv[1]))]
    return "\n".join(parts), "; ".join(unexplained[:5])


def _unnamed_addresses(session: Session) -> dict[str, list[str]]:
    """Addresses held only under placeholder names, per domain (evidence for Tier E)."""
    out: dict[str, set[str]] = defaultdict(set)
    for email, domain, name in session.execute(
            select(EmailObservation.email, EmailObservation.domain, EmailObservation.first_name_raw)
            .where(EmailObservation.excluded_reason == "placeholder_name")):
        out[domain].add(email)
    return {d: sorted(v) for d, v in out.items()}


def build(session: Session, run: InferenceRun, path: Path | None = None) -> Path:
    path = Path(path or config.DATA_OUT / "firm_tracker.xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)

    views, _ = pipeline._views(session)
    samples = inference.build_samples(views, use_nicknames=run.use_nicknames)
    unnamed = _unnamed_addresses(session)
    dps = session.scalars(select(DomainPattern).where(DomainPattern.run_id == run.id)).all()
    dps.sort(key=lambda d: (PRIORITY[d.tier], -d.contacts_held, d.domain))

    wb = Workbook()
    ws = wb.active
    ws.title = "Firms"
    cols = ["Firm (email domain)", "Tier", "Status", "Action needed", "Current format", "Example",
            "Agreement", "Addresses explained", "Contacts held", "Evidence by format",
            "Addresses we can't explain", "Addresses with no name",
            "Correct format", "Example correct address", "Checked by", "Notes"]
    widths = [28, 6, 22, 44, 14, 34, 10, 11, 9, 70, 44, 44, 22, 34, 14, 36]
    manual = cols[-4:]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"

    for dp in dps:
        status, action = STATUS[dp.tier]
        if dp.migrated_from:
            status, action = "Format changed recently", f"Was {dp.migrated_from}; confirm the newer format is current."
        by_format, unexplained = _evidence(samples.get(dp.domain, []))
        nameless = unnamed.get(dp.domain, [])
        ws.append([
            dp.domain, dp.tier, status, action,
            dp.pattern if dp.tier != "E" else None,
            worked_example(dp.pattern, dp.domain) if dp.tier in "ABC" else None,
            dp.dominance, dp.sample_count - dp.unexplained_count, dp.contacts_held,
            by_format or None, unexplained or None,
            ("; ".join(nameless[:5]) + (f" (+{len(nameless) - 5})" if len(nameless) > 5 else "")) or None,
            None, None, None, None,
        ])
        r = ws.max_row
        ws.cell(row=r, column=2).fill = TIER_FILLS[dp.tier]
        ws.cell(row=r, column=10).alignment = Alignment(wrap_text=True, vertical="top")
        for j in range(len(cols) - len(manual) + 1, len(cols) + 1):
            ws.cell(row=r, column=j).fill = INPUT_FILL

    last = max(ws.max_row, 2)
    # Dropdown source lives on a hidden sheet: too many choices for an inline list.
    lists = wb.create_sheet("Formats")
    choices = [f"{p.name}  ({p.example()})" for p in library.PATTERNS] + list(EXTRA_CHOICES)
    for c in choices:
        lists.append([c])
    lists.sheet_state = "hidden"
    fmt_col = get_column_letter(cols.index("Correct format") + 1)
    dv = DataValidation(type="list", formula1=f"=Formats!$A$1:$A${len(choices)}", allow_blank=True)
    dv.add(f"{fmt_col}2:{fmt_col}{last}")
    ws.add_data_validation(dv)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    about = wb.create_sheet("About", 0)
    counts = {t: sum(1 for d in dps if d.tier == t) for t in STATUS}
    for row in [
        ("Firm tracker", ""),
        ("What this is", "Every firm we hold contacts for, worst problems first. Yellow columns are for "
                         "whoever checks the firm. Filter the Tier column to work one group at a time."),
        ("Tier D", f"{counts['D']} firms: our addresses disagree on the format. Evidence shows which named "
                   "people use which format."),
        ("Tier E", f"{counts['E']} firms: every address we hold has no usable first name, so no format "
                   "can be learned. One known name + address fixes it."),
        ("Tier C", f"{counts['C']} firms: format learned from a single address."),
        ("Tier B / A", f"{counts['B']} / {counts['A']} firms: two agreeing / three or more agreeing addresses."),
        ("Agreement", "Share of the firm's address evidence that fits the current format (1 = all)."),
        ("Inference run", f"#{run.id}, nicknames {'on' if run.use_nicknames else 'off'}"),
    ]:
        about.append(list(row))
    about["A1"].font = Font(bold=True, size=13)
    about.column_dimensions["A"].width = 16
    about.column_dimensions["B"].width = 110
    for r in about.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.active = 1

    wb.save(path)
    return path
