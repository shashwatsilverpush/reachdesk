"""Validation capture sheet: data/out/validation_sheet.xlsx.

A seeded, stratified random sample of domains from Tiers A, B and C. Computed
columns are filled; the yellow columns are for manual entry while checking
predictions against an external source.
"""

from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from leadcore import config
from leadcore.models import DomainPattern, InferenceRun
from leadcore.report.workbook import HEADER_FILL, HEADER_FONT, worked_example

INPUT_FILL = PatternFill("solid", start_color="FFFF99")
COMPUTED = ["Tier", "Domain", "Samples", "Inferred pattern", "Worked example", "Evidence address"]
MANUAL = ["Target first name", "Target last name", "Our predicted address",
          "Externally revealed address", "Verdict", "Notes"]
VERDICTS = ("match", "mismatch", "no-data", "skipped")


def sample_domains(session: Session, run: InferenceRun, per_tier: int, seed: int) -> list[DomainPattern]:
    rng = random.Random(seed)
    picked: list[DomainPattern] = []
    for tier in config.GENERATABLE_TIERS:
        pool = session.scalars(select(DomainPattern).where(DomainPattern.run_id == run.id,
                                                           DomainPattern.tier == tier)
                               .order_by(DomainPattern.domain)).all()
        picked.extend(sorted(rng.sample(pool, min(per_tier, len(pool))), key=lambda d: d.domain))
    return picked


def build(session: Session, run: InferenceRun, *, per_tier: int | None = None, seed: int | None = None,
          path: Path | None = None) -> Path:
    per_tier = per_tier or config.CAPTURE_SHEET_PER_TIER
    seed = config.CAPTURE_SHEET_SEED if seed is None else seed
    path = Path(path or config.DATA_OUT / "validation_sheet.xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Validation"
    cols = COMPUTED + MANUAL
    widths = [6, 30, 9, 16, 38, 40, 18, 18, 38, 38, 12, 40]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    rows = sample_domains(session, run, per_tier, seed)
    for dp in rows:
        ws.append([dp.tier, dp.domain, dp.sample_count, dp.pattern,
                   worked_example(dp.pattern, dp.domain), dp.evidence_email] + [None] * len(MANUAL))
        for j in range(len(COMPUTED) + 1, len(cols) + 1):
            ws.cell(row=ws.max_row, column=j).fill = INPUT_FILL

    last = max(ws.max_row, 2)
    verdict_col = get_column_letter(cols.index("Verdict") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Verdict", error="Choose " + " / ".join(VERDICTS))
    dv.add(f"{verdict_col}2:{verdict_col}{last}")
    ws.add_data_validation(dv)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    info = wb.create_sheet("About")
    for line in (
        ("Validation capture sheet", ""),
        ("Inference run", run.id),
        ("Nicknames", "on" if run.use_nicknames else "off"),
        ("Per tier", per_tier),
        ("Seed", seed),
        ("How to use", "For each domain, pick a real person there, enter their name, compute our "
                       "predicted address from the inferred pattern, look up the address from an "
                       "external source, and record the verdict."),
    ):
        info.append(list(line))
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 100

    wb.save(path)
    return path
